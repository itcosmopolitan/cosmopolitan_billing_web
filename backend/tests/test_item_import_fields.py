from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

from src.database import Base
from src.models import Branch, Item, ItemBranchConfig, User
from src.routes.items import download_import_template, process_item_import


@pytest.mark.asyncio
async def test_item_import_persists_packing_and_fixed_rates():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    session_factory = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "Name", "Packing", "Wholesale Rate", "Staff Rate",
        "Opening Stock - Main", "Cost Price - Main", "Selling Price - Main",
    ])
    sheet.append(["Imported item", "Box of 6", 80.5, 65, None, 10, 20])
    sheet.append(["Zero stock item", None, None, None, 0, 11, 21])
    content = BytesIO()
    workbook.save(content)

    async with session_factory() as db:
        user = User(
            id="import-user",
            name="Import user",
            hashed_password="unused",
            all_branches=True,
        )
        db.add_all([user, Branch(id="branch-main", name="Main", code="MAIN")])
        await db.commit()

        result = await process_item_import(content.getvalue(), db, user)
        imported = (await db.execute(select(Item).where(Item.name == "Imported item"))).scalar_one()
        zero_stock_item = (
            await db.execute(select(Item).where(Item.name == "Zero stock item"))
        ).scalar_one()
        branch_configs = (
            await db.execute(select(ItemBranchConfig).order_by(ItemBranchConfig.item_id))
        ).scalars().all()

        assert result["created"] == 2
        assert imported.packaging == "Box of 6"
        assert imported.is_packaging is True
        assert imported.wholesale_pricing_mode == "price"
        assert imported.wholesale_price == 80.5
        assert imported.staff_pricing_mode == "price"
        assert imported.staff_price == 65
        assert len(branch_configs) == 1
        assert branch_configs[0].item_id == zero_stock_item.id
        assert branch_configs[0].is_available is True

        response = await download_import_template(db)
        template_content = b"".join([chunk async for chunk in response.body_iterator])
        template = openpyxl.load_workbook(BytesIO(template_content), read_only=True)
        headers = next(template.active.iter_rows(values_only=True))
        assert {"Packing", "Wholesale Rate", "Staff Rate"}.issubset(headers)
        template.close()

    await engine.dispose()